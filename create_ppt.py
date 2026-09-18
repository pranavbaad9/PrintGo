from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Helper function to add slides
def add_slide(title_text, bullet_points, notes_text, layout_index=1):
    slide_layout = prs.slide_layouts[layout_index] # 0 is title, 1 is title and content
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = title_text
    
    if layout_index == 1:
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        
        for i, point in enumerate(bullet_points):
            if i == 0:
                p = tf.paragraphs[0]
                p.text = point
            else:
                p = tf.add_paragraph()
                p.text = point
                p.level = 0
                
    # Add notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text
    
    return slide

# Slide 1: Title
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "PrintGo"
subtitle.text = "A smart, unattended printing platform connecting customers, kiosks, cloud services, and physical printers."
slide.notes_slide.notes_text_frame.text = "Welcome everyone. Today I'm excited to introduce PrintGo, a revolutionary platform that brings printing into the modern, self-service era."

# Slide 2: The Problem
add_slide(
    "The Problem",
    [
        "Traditional print shops require handing personal documents to operators, compromising privacy.",
        "Customers waste valuable time waiting in lines.",
        "Shops are limited by operating hours and staff availability.",
        "Manual payment and refund processes are tedious and error-prone."
    ],
    "Speaker Notes: Have you ever needed to print a confidential document and felt uneasy handing it over to a shop operator on a USB drive? Or waited in line for 20 minutes just to print a single page? Traditional print shops are fundamentally broken for the modern user. They are slow, lack privacy, have limited hours, and the payment process is entirely manual."
)

# Slide 3: The Solution
add_slide(
    "The Solution",
    [
        "Faster, self-service printing via mobile devices.",
        "Strict document privacy (files are auto-deleted 5 minutes after printing).",
        "24/7 unattended operation through standalone kiosks.",
        "Fully automated payment and refund workflows if hardware fails."
    ],
    "Speaker Notes: Enter PrintGo. We completely remove the middleman. Customers control the entire process from their phones. We guarantee privacy by automatically wiping files 5 minutes after printing. Our kiosks operate 24/7 without staffing, and we have fully automated the trickiest part of self-service hardware: payments and refunds."
)

# Slide 4: How it Works (User Flow)
add_slide(
    "How It Works (User Flow)",
    [
        "1. Scan: Customer scans a QR code on the PrintGo kiosk with their mobile device.",
        "2. Upload: Customer uploads their document directly from their phone.",
        "3. Configure: Customer selects print settings (color, copies, etc.).",
        "4. Pay: Customer pays seamlessly via Cashfree.",
        "5. Print: The Kiosk instantly prints the document."
    ],
    "Speaker Notes: The user experience is frictionless. You walk up to a kiosk, scan a QR code to link your phone, upload your file, choose your settings, and pay. The kiosk instantly wakes up and prints your document. It’s entirely touchless for the kiosk screen until you grab your paper."
)

# Slide 5: Architecture
add_slide(
    "Architecture",
    [
        "Modern, real-time distributed system.",
        "Cloud Backend: Handles core business logic, API, and state.",
        "Database: PostgreSQL for reliable relational data storage.",
        "Kiosk & Mobile UIs: Connected via real-time WebSocket sessions.",
        "Printer Agent: Bridges the cloud with physical hardware."
    ],
    "Speaker Notes: Under the hood, PrintGo relies on a modern, real-time architecture. Our Cloud Backend orchestrates everything. It talks to PostgreSQL for data, Cashfree for payments, and maintains persistent WebSocket connections with both the customer's mobile device and the physical kiosk."
)

# Slide 6: Tech Stack
add_slide(
    "Tech Stack",
    [
        "Backend: Node.js, Express, Socket.io (real-time sync).",
        "Database: PostgreSQL with Prisma ORM.",
        "Frontend: React and Vite (Kiosk display & Admin dashboards).",
        "Agent: Windows-based Node.js desktop app (polls Windows spooler).",
        "Payments: Cashfree integration."
    ],
    "Speaker Notes: We chose a robust Javascript-based stack for rapid iteration and real-time capabilities. Node.js and Socket.io power our backend. React drives our user interfaces. Crucially, our Printer Agent is a custom Windows app that deeply integrates with the OS print spooler to monitor physical printer health."
)

# Slide 7: Security & Privacy
add_slide(
    "Security & Privacy",
    [
        "Privacy First: Documents securely wiped 5 mins after job completion.",
        "Secure Auth: JWT for web authentication, static Machine Keys for kiosks.",
        "Safe Inputs: Strict validation to prevent command injection at the OS level.",
        "Data Lifecycle: Upload -> Process -> Store -> Print -> Cleanup -> Delete."
    ],
    "Speaker Notes: Security and privacy are our core pillars. Your files are never stored permanently; they are wiped just 5 minutes after printing. We use strict authentication protocols, ensuring kiosks can only execute authorized print jobs, preventing malicious OS-level attacks."
)

# Slide 8: Business Model & Scalability
add_slide(
    "Business Model & Scalability",
    [
        "Multi-Tenant Design: Built for scale with Franchisee and SuperAdmin roles.",
        "Tenant Isolation: Franchisees can only access and manage their own kiosks and data.",
        "Automated Management: Over-The-Air (OTA) software updates for kiosks.",
        "Hardware-Aware Refunds: Agent detects paper jams and triggers automatic API refunds."
    ],
    "Speaker Notes: PrintGo is designed to scale via a Franchisee model. Our multi-tenant architecture ensures strict data isolation between franchisees. Furthermore, we drastically reduce operational overhead by pushing Over-The-Air updates to the kiosks and handling hardware-failure refunds completely automatically."
)

# Slide 9: Roadmap
add_slide(
    "Roadmap & Future Vision",
    [
        "Current: Solidifying payments and spooler polling accuracy.",
        "Next Phase: Implementing Redis for horizontal cloud scaling; Razorpay integration.",
        "Future: Highly optimized native C++/Rust Printer Agent.",
        "Future: Hardware-level encryption and international payment gateways."
    ],
    "Speaker Notes: We have a clear path forward. Currently, we are hardening our payment flows. Next, we will introduce Redis to horizontally scale our cloud infrastructure and add Razorpay. Long term, we plan to rebuild our Printer Agent in Rust or C++ for a lighter footprint and introduce hardware-level encryption."
)

# Slide 10: Thank You
add_slide(
    "Thank You",
    [
        "PrintGo",
        "Smart, Unattended Printing",
        "Any Questions?"
    ],
    "Speaker Notes: Thank you for your time. I'd be happy to answer any questions about the PrintGo platform, our technology, or our business model."
)

prs.save('c:/Projects/PrintGo/PrintGo_Presentation.pptx')
